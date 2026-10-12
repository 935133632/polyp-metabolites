"""Small python-docx helpers for journal-style output: three-line tables, inline **bold** / *italic* markup, figures with legends."""
import re
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def new_document(font="Times New Roman", size=11, line_spacing=1.5):
    d = Document()
    st = d.styles["Normal"]; st.font.name = font; st.font.size = Pt(size)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    st.paragraph_format.line_spacing = line_spacing; st.paragraph_format.space_after = Pt(6)
    for s in d.sections: s.left_margin = s.right_margin = Cm(2.5); s.top_margin = s.bottom_margin = Cm(2.5)
    for name, sz in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        h = d.styles[name]; h.font.name = font; h.font.size = Pt(sz); h.font.bold = True; h.font.color.rgb = None
        h.element.rPr.rFonts.set(qn("w:eastAsia"), font); h.paragraph_format.space_before = Pt(12)
    return d


def add_runs(par, text, size=None):
    """Write text with **bold**, *italic* and ^superscript^ markup into a paragraph."""
    for tok in re.split(r"(\*\*.+?\*\*|\*.+?\*|\^.+?\^)", text):
        if not tok: continue
        if tok.startswith("**"): r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*") and len(tok) > 2: r = par.add_run(tok[1:-1]); r.italic = True
        elif tok.startswith("^") and len(tok) > 2: r = par.add_run(tok[1:-1]); r.font.superscript = True
        else: r = par.add_run(tok)
        if size: r.font.size = Pt(size)
    return par


def _border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr(); b = tcPr.find(qn("w:tcBorders"))
    if b is None: b = OxmlElement("w:tcBorders"); tcPr.append(b)
    for edge, (sz, val) in kw.items():
        e = b.find(qn(f"w:{edge}"))
        if e is None: e = OxmlElement(f"w:{edge}"); b.append(e)
        e.set(qn("w:val"), val); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")


def landscape(section):
    section.orientation = WD_ORIENT.LANDSCAPE; w, h = section.page_width, section.page_height
    if w < h: section.page_width, section.page_height = h, w
    section.left_margin = section.right_margin = Cm(1.8)


def three_line_table(doc, header_rows, body_rows, col_widths_cm=None, font_size=8, merges=(), bold_cells=(), indent_rows=(), group_rows=()):
    """header_rows / body_rows: lists of lists of strings. merges: [(row, col0, col1)] over header rows. bold_cells: {(body_row, col)}.
    indent_rows: body rows to indent in column 0; group_rows: body rows rendered as bold group labels."""
    nh = len(header_rows); t = doc.add_table(rows=nh + len(body_rows), cols=len(header_rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i, row in enumerate(header_rows + body_rows):
        for j, txt in enumerate(row):
            c = t.cell(i, j); c.text = ""; p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            add_runs(p, str(txt), size=font_size)
            if i < nh or j > 0: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i < nh or (i - nh) in group_rows or (i - nh, j) in set(bold_cells):
                for r in p.runs: r.bold = True
            if j == 0 and i >= nh and (i - nh) in indent_rows: p.paragraph_format.left_indent = Cm(0.4)
            if col_widths_cm: c.width = Cm(col_widths_cm[j])
    for (r, c0, c1) in merges: t.cell(r, c0).merge(t.cell(r, c1))
    for j in range(len(header_rows[0])):
        _border(t.cell(0, j), top=(12, "single")); _border(t.cell(nh - 1, j), bottom=(6, "single")); _border(t.cell(nh + len(body_rows) - 1, j), bottom=(12, "single"))
    return t


def caption(doc, text, size=10, bold_prefix=True):
    p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.15
    add_runs(p, text, size=size); return p


def add_figure(doc, path, width_cm, legend, size=9):
    doc.add_picture(str(path), width=Cm(width_cm)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.15; add_runs(p, legend, size=size); return p
